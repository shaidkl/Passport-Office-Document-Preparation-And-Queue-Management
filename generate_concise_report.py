from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / ".codex_deps"))

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = ROOT / "Recommendation Report - Supporting Women Entrepreneurs in Nepal.docx"
NAVY = "17365D"
TEAL = "0E7490"
PALE_TEAL = "E6F4F1"
LIGHT = "F2F4F7"
WHITE = "FFFFFF"
GREY = "667085"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), color)


def cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = tc_mar.find(qn(f"w:{name}"))
        if item is None:
            item = OxmlElement(f"w:{name}")
            tc_mar.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    item = OxmlElement("w:tblHeader")
    item.set(qn("w:val"), "true")
    tr_pr.append(item)


def style_table(table, font_size=9, first_col_bold=False):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(table.rows):
        no_split(row)
        if r_idx == 0:
            repeat_header(row)
        for c_idx, cell in enumerate(row.cells):
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0:
                shade(cell, NAVY)
            elif r_idx % 2 == 0:
                shade(cell, LIGHT)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    run.font.size = Pt(font_size)
                    if r_idx == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string(WHITE)
                    elif first_col_bold and c_idx == 0:
                        run.bold = True


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def hyperlink(paragraph, label, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = label
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def paragraph(document, text="", style=None, lead=None):
    p = document.add_paragraph(style=style)
    if lead and text.startswith(lead):
        p.add_run(lead).bold = True
        p.add_run(text[len(lead):])
    else:
        p.add_run(text)
    return p


def bullet(document, text, level=0):
    return paragraph(document, text, "List Bullet" if level == 0 else "List Bullet 2")


def numbered(document, text):
    return paragraph(document, text, "List Number")


def callout(document, title, text):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, PALE_TEAL)
    cell_margins(cell, 150, 180, 150, 180)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)


def contents(document):
    entries = [
        "Table of Contents",
        "Executive Summary",
        "1. Introduction",
        "2. Methodology",
        "3. Main Findings",
        "3.1 Women participate, but most enterprises remain small",
        "3.2 Finance is available but difficult to use",
        "3.3 Training is not connected to business growth",
        "3.4 Formalisation and market access are difficult",
        "3.5 Care work and social norms limit opportunity",
        "4. Recommendations and Discussion",
        "4.1 Establish local women’s enterprise desks",
        "4.2 Create a finance ladder for business growth",
        "4.3 Provide mentoring instead of one-time training",
        "4.4 Connect women to repeat buyers and procurement",
        "4.5 Address care, safety and accountability",
        "5. Next Steps and Conclusion",
        "Appendix A: Action Plan",
        "Appendix B: Presentation Guide",
        "References",
    ]
    table = document.add_table(rows=0, cols=3)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for title in entries:
        cells = table.add_row().cells
        cells[0].width = Cm(13.4)
        cells[1].width = Cm(2.4)
        cells[2].width = Cm(1.0)
        for cell in cells:
            cell_margins(cell, 16, 0, 16, 0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cells[0].text = title
        cells[1].text = "........................"
        cells[2].text = "–"
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for cell in cells:
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9.4)
        if title.startswith(("Table", "Executive", "1.", "2.", "3. Main", "4. Recommendations", "5.", "Appendix", "References")) and not title.startswith(("3.1", "3.2", "3.3", "3.4", "3.5", "4.1", "4.2", "4.3", "4.4", "4.5")):
            cells[0].paragraphs[0].runs[0].bold = True
    return table


def reference(document, citation, url):
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(citation + " ")
    hyperlink(p, "Available online", url)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.0)
section.header_distance = Cm(0.7)
section.footer_distance = Cm(0.7)
section.different_first_page_header_footer = True

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string("263238")
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(4.5)

for name in ("List Bullet", "List Bullet 2", "List Number"):
    doc.styles[name].font.name = "Aptos"
    doc.styles[name].font.size = Pt(10.3)
    doc.styles[name].paragraph_format.space_after = Pt(2.5)

for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 12.5, TEAL)):
    style = doc.styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.text = "CONCISE RECOMMENDATION REPORT   |   NEPAL"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(GREY)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Supporting Women Entrepreneurs in Nepal   •   ")
add_page_number(footer)
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GREY)

# Title page
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(70)
r = p.add_run("RECOMMENDATION REPORT")
r.font.size = Pt(12)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(TEAL)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run("Supporting Women\nEntrepreneurs in Nepal")
r.font.name = "Aptos Display"
r.font.size = Pt(28)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(NAVY)

p = doc.add_paragraph()
r = p.add_run("A concise report designed for classroom presentation")
r.font.size = Pt(14)
r.italic = True
r.font.color.rgb = RGBColor.from_string(GREY)

doc.add_paragraph().paragraph_format.space_after = Pt(30)
callout(
    doc,
    "Main message",
    "Nepal should connect registration, finance, mentoring, markets and practical support into one clear pathway that helps women move from very small businesses to stable, growing enterprises.",
)
doc.add_paragraph().paragraph_format.space_after = Pt(45)

info = doc.add_table(rows=3, cols=2)
for i, (label, value) in enumerate(
    (("Prepared by", "________________________"), ("Course / institution", "________________________"), ("Date", "September 2026"))
):
    info.cell(i, 0).text = label
    info.cell(i, 1).text = value
    info.cell(i, 0).paragraphs[0].runs[0].bold = True
    for cell in info.rows[i].cells:
        cell_margins(cell, 50, 0, 50, 100)

doc.add_page_break()

doc.add_heading("Table of Contents", 1)
contents(doc)
doc.add_page_break()

# Executive Summary
doc.add_heading("Executive Summary", 1)
paragraph(
    doc,
    "Women play an important role in Nepal’s business economy, but most women-led enterprises remain very small. The National Economic Census 2018 found that women owned 26.9 per cent of establishments and managed 29.6 per cent. However, 98.3 per cent of female-managed establishments were micro enterprises with fewer than ten workers. Their average sales and profits were also much lower than those of male-managed businesses.",
)
paragraph(
    doc,
    "Nepal already offers useful support, including registration-fee discounts, concessional loans and financial-inclusion policies. The main problem is that these services are scattered. Women may receive training but no market link, or register a business but still fail to obtain suitable finance. Care duties, limited mobility, lack of property for collateral and weak business networks make growth harder.",
)
paragraph(doc, "This report recommends five connected actions:")
for item in [
    "Create women’s enterprise service desks in selected municipalities.",
    "Offer finance suited to start-up, working-capital and growth stages.",
    "Replace one-time training with practical mentoring for up to twelve months.",
    "Connect women-owned businesses to regular buyers and public procurement.",
    "Include childcare, safety, local access and transparent monitoring in every programme.",
]:
    bullet(doc, item)

# Introduction
doc.add_heading("1. Introduction", 1)
paragraph(
    doc,
    "Women across Nepal run farms, shops, food businesses, tourism services, handicraft enterprises and new digital ventures. These activities support families and local economies. Still, many businesses remain informal and earn only a small profit. Supporting women entrepreneurs is therefore not only a gender-equality issue; it is also a way to increase income, create jobs and reduce economic dependence.",
)
paragraph(
    doc,
    "The World Bank estimated women’s paid labour-force participation at 24.4 per cent in 2023/24, compared with 52.7 per cent for men. Entrepreneurship can create opportunities where salaried jobs are limited, especially outside large cities. However, simply counting a woman as self-employed is not enough. A good policy should help her control the business, earn a stable income and expand when there is market demand.",
)
callout(
    doc,
    "Presentation takeaway",
    "The issue is not whether Nepali women are entrepreneurial. The issue is whether their enterprises receive enough connected support to survive and grow.",
)

# Methodology
doc.add_heading("2. Methodology", 1)
paragraph(
    doc,
    "This report is based on a review of Nepal’s laws, official statistics and recent reports from the National Statistics Office, Ministry of Finance, Nepal Rastra Bank, ILO, UN Women and the World Bank. The evidence was grouped into five themes: enterprise size, finance, skills, markets, and social or institutional barriers. No primary interviews were conducted. Because the latest complete Economic Census data are from 2018, the report recommends new local-level data and pilot programmes before nationwide expansion.",
)

# Findings
doc.add_heading("3. Main Findings", 1)
doc.add_heading("3.1 Women participate, but most enterprises remain small", 2)
paragraph(
    doc,
    "The Economic Census counted 247,880 women-owned establishments. This shows strong participation, but women were concentrated in micro businesses and lower-capital sectors such as retail, food, accommodation and personal services. Female-managed establishments recorded average annual sales of about NPR 1 million, compared with NPR 4.1 million for male-managed establishments. The policy challenge is therefore growth and productivity, not only business creation.",
)

doc.add_heading("3.2 Finance is available but difficult to use", 2)
paragraph(
    doc,
    "The Economic Survey 2023/24 reported NPR 46.48 billion in outstanding women-entrepreneur concessional loans for 76,975 borrowers in mid-March 2024. Yet many women lack property in their own names, formal accounts or a long credit history. Microfinance may be easier to access but often provides loans too small or too short for expansion. Women also need clearer information about eligibility, costs and repayment.",
)

doc.add_heading("3.3 Training is not connected to business growth", 2)
paragraph(
    doc,
    "Short courses can increase knowledge, but they rarely solve problems such as pricing, packaging, taxation, quality standards, digital security and buyer negotiation. The ILO’s study of Kathmandu and Pokhara found gaps in finance, human capital, markets, policy, culture and support. Women need an adviser who follows their progress and connects learning to finance and real customers.",
)

doc.add_heading("3.4 Formalisation and market access are difficult", 2)
paragraph(
    doc,
    "Around half of the establishments counted in 2018 were unregistered. Registration may require visits to local offices, tax authorities and industry offices. Women are unlikely to formalise if the process is confusing and registration brings no clear benefit. At the same time, small production volume, weak transport links and limited networks prevent many women from reaching larger or repeat buyers.",
)

doc.add_heading("3.5 Care work and social norms limit opportunity", 2)
paragraph(
    doc,
    "Many women manage a business while also carrying most childcare, elder care and household work. This reduces the time available for production, travel, training and banking. Safety concerns and expectations about women’s movement can also restrict business choices. These are economic barriers because they directly affect working hours, market access and the ability to take business risks.",
)

doc.add_heading("Key Evidence at a Glance", 2)
stats = [
    ("26.9%", "Establishments with female owners"),
    ("29.6%", "Establishments with female managers"),
    ("98.3%", "Female-managed establishments that were micro-sized"),
    ("NPR 46.48bn", "Outstanding women-entrepreneur loans in mid-March 2024"),
    ("24.4%", "Women’s paid labour-force participation in 2023/24"),
    ("0.53", "Nepal’s Financial Inclusion Index in 2024/25"),
]
table = doc.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = "Figure"
table.rows[0].cells[1].text = "Meaning"
for figure, meaning in stats:
    cells = table.add_row().cells
    cells[0].text = figure
    cells[1].text = meaning
style_table(table, 9.2, True)

# Recommendations
doc.add_heading("4. Recommendations and Discussion", 1)
callout(doc, "Recommended approach", "Create one simple pathway: register → finance → improve → sell → grow.")

doc.add_heading("4.1 Establish local women’s enterprise desks", 2)
paragraph(
    doc,
    "Pilot a single support desk in selected urban and rural municipalities. A trained adviser should explain registration and tax requirements, help prepare documents and refer women to banks, mentors, certification services and buyers. Mobile clinics and local-language support are necessary for remote communities.",
)
bullet(doc, "Lead: Ministry of Industry, provincial and local governments.")
bullet(doc, "Measure: completed registrations, referral use, processing time and cost.")

doc.add_heading("4.2 Create a finance ladder for business growth", 2)
paragraph(
    doc,
    "Provide different products for start-ups, working capital and business expansion. Banks should assess cash flow, digital payments, purchase orders and movable assets instead of relying mainly on land. Partial credit guarantees can reduce lender risk. Rejection reasons, approval time, loan size and province-level reach should be published.",
)
bullet(doc, "Lead: Nepal Rastra Bank, Deposit and Credit Guarantee Fund and banks.")
bullet(doc, "Measure: approval rate, repayment, average loan size and business growth.")

doc.add_heading("4.3 Provide mentoring instead of one-time training", 2)
paragraph(
    doc,
    "Give selected entrepreneurs practical support for up to twelve months. Monthly mentoring should cover costing, cash flow, digital payments, quality, marketing, tax and customer service. Programmes should group women by business stage and sector. Providers should be assessed by business improvements rather than the number of trainees.",
)
bullet(doc, "Lead: Provincial and local governments with CTEVT, universities and business groups.")
bullet(doc, "Measure: accounts maintained, repeat customers, sales and standards achieved.")

doc.add_heading("4.4 Connect women to repeat buyers and procurement", 2)
paragraph(
    doc,
    "Create a verified directory of women-owned and controlled businesses. Public agencies and large private companies should publish purchasing opportunities, divide suitable contracts into smaller lots and run supplier-development programmes. Producer groups can combine volume, transport and packaging. The aim is regular orders, not only one-day exhibitions.",
)
bullet(doc, "Lead: Industry ministry, procurement authorities, chambers and FWEAN.")
bullet(doc, "Measure: value of contracts, repeat orders and provincial distribution.")

doc.add_heading("4.5 Address care, safety and accountability", 2)
paragraph(
    doc,
    "Enterprise programmes should include local scheduling, childcare support, safe travel and a confidential complaint process. A national coordination group should adopt one definition of a women-owned and controlled enterprise and publish annual results. Data should show who applied, who was rejected and whether supported businesses survived, increased sales or created decent jobs.",
)
bullet(doc, "Lead: All programme owners, coordinated by the industry ministry.")
bullet(doc, "Measure: participation across social groups, business survival and women’s control over decisions.")

# Conclusion
doc.add_heading("5. Next Steps and Conclusion", 1)
paragraph(doc, "The recommended actions can be introduced in stages:")
numbered(doc, "First six months: agree a common definition, select pilot municipalities and map existing services.")
numbered(doc, "Months six to twelve: open local desks, train bank focal persons and begin mentoring and supplier programmes.")
numbered(doc, "At twelve months: publish the first results dashboard and correct delays or exclusion.")
numbered(doc, "At twenty-four months: independently review business survival, sales, jobs and women’s decision-making, then scale what works.")
paragraph(
    doc,
    "Nepal already has active women entrepreneurs, supportive laws, credit schemes and business networks. The missing element is a clear connection between these resources. A woman should not have to move from office to office or repeat beginner training while still lacking finance and customers. Local entry, suitable credit, long-term mentoring, dependable markets and practical care support would make current spending more effective.",
)
paragraph(
    doc,
    "Success should mean more than registering businesses or training participants. It should mean that women genuinely control productive enterprises, earn stable incomes, create safe jobs and take part in economic decisions. A focused pilot-and-learn approach gives Nepal a realistic way to achieve that result.",
)

# Appendices
doc.add_page_break()
doc.add_heading("Appendix A: Action Plan", 1)
actions = [
    ("Local service desks", "0–9 months", "Registration, referrals, time and cost"),
    ("Finance ladder", "3–12 months", "Approvals, repayment, loan size and reach"),
    ("Twelve-month mentoring", "6–18 months", "Sales, records, standards and customers"),
    ("Supplier development", "6–24 months", "Contracts and repeat orders"),
    ("Shared dashboard", "From month 6", "Survival, jobs and women’s decision-making"),
]
table = doc.add_table(rows=1, cols=3)
for i, title in enumerate(("Action", "Timing", "Main measure")):
    table.rows[0].cells[i].text = title
for row in actions:
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
style_table(table, 9, True)

doc.add_heading("Appendix B: Presentation Guide", 1)
slides = [
    "Slide 1 — Topic and central message",
    "Slide 2 — Why women’s entrepreneurship matters in Nepal",
    "Slide 3 — Key statistics",
    "Slide 4 — Main barriers: finance, skills, markets and care",
    "Slide 5 — Recommendations 1 and 2: local desks and finance",
    "Slide 6 — Recommendations 3 and 4: mentoring and buyers",
    "Slide 7 — Recommendation 5 and implementation timeline",
    "Slide 8 — Conclusion: measure business growth, not attendance",
]
for item in slides:
    bullet(doc, item)

# References
doc.add_page_break()
doc.add_heading("References", 1)
refs = [
    ("Central Bureau of Statistics. (2021). National Economic Census 2018: Analytical report—Women in business.", "https://giwmscdnone.gov.np/media/app/public/36/posts/1693998585_42.pdf"),
    ("Department of Industry. (2020). The Industrial Enterprises Act, 2076 (2020).", "https://www.doind.gov.np/detail/5c6a4832-321c-40b9-8b54-a0f4e94edce0"),
    ("International Labour Organization. (2023). Building inclusive entrepreneurship ecosystems in Nepal.", "https://www.ilo.org/publications/building-inclusive-entrepreneurship-ecosystems-nepal-analysis-kathmandu-and"),
    ("Ministry of Finance. (2024). Economic Survey 2023/24.", "https://giwmscdntwo.gov.np/media/pdf_upload/MOF_Economic%20Survey%20ENG%202023-24%20book%20Final_for%20WEB_miwtpw0.pdf"),
    ("Nepal Rastra Bank. (2026). Financial Inclusion Index for Nepal 2025.", "https://www.nrb.org.np/contents/uploads/2026/02/FInancial-Inclusion-Index-2025.pdf"),
    ("UN Women. (2023). Accelerating the financial inclusion of women in Nepal.", "https://asiapacific.unwomen.org/en/digital-library/publications/2023/07/accelerating-the-financial-inclusion-of-women-in-nepal"),
    ("World Bank. (2024). Women’s labor force participation in Nepal: An exploration of the role of social norms.", "https://documents.worldbank.org/en/publication/documents-reports/documentdetail/099816406182414601"),
    ("World Bank. (2025). Nurturing Nepali talent to foster economic growth.", "https://thedocs.worldbank.org/en/doc/046021b791dda7b320bc3b778fcd99b6-0310012025/original/Nepal-HCR-long-PPT-Feb-28-2025.pdf"),
]
for citation, url in refs:
    reference(doc, citation, url)

doc.core_properties.title = "Recommendation Report on Supporting Women Entrepreneurs in Nepal"
doc.core_properties.subject = "Concise, presentation-friendly recommendation report"
doc.core_properties.keywords = "Nepal, women entrepreneurs, recommendation report, presentation"

for p in doc.paragraphs:
    p.paragraph_format.widow_control = True

doc.save(OUTPUT)

text_parts = []
started = False
for p in doc.paragraphs:
    if p.text == "Executive Summary":
        started = True
    if p.text == "Appendix A: Action Plan":
        started = False
    if started:
        text_parts.append(p.text)
word_count = len(re.findall(r"\b[\w’'-]+\b", " ".join(text_parts)))
print(f"Created: {OUTPUT}")
print(f"Approximate main-report word count: {word_count}")
